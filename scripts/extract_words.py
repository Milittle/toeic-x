#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl"]
# ///
"""Extract the TOEIC vocabulary workbook into a static JSON library.

Source : ``单词本.xlsx`` (sheets 重点1500 / TOEIC专项_TSL1250 / 基础_NGSL2809)
Output : ``data/words/words.json``  (read-only, ADR-0005)

Words are a static reference corpus parallel to the question bank (ADR-0002)
and the collocations library (ADR-0004) — read-only, regenerated manually via
``uv run scripts/extract_words.py``. The app never writes this file.

Each row is tagged with its word list (``list``); the same headword may appear
in several lists (e.g. *refund* is in both 重点1500 and TSL1250) and is kept as
separate rows — the vocab page browses by list, it does not dedupe.

The 关联固定搭配 cell is matched against ``data/collocations/collocations.json``
by expression text to produce a clickable ``collocationId``; unmatched labels
are left as plain text (no link). Question reachability is mediated by the
collocations library, which already gates unseen-question samples — the vocab
page never links a word directly to a question number (ADR-0005).
"""

from __future__ import annotations

import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

# sheet name -> list key + human label
SHEETS: list[tuple[str, str, str]] = [
    ("重点1500", "key1500", "重点1500"),
    ("TOEIC专项_TSL1250", "tsl1250", "TOEIC专项 TSL1250"),
    ("基础_NGSL2809", "ngsl2809", "基础 NGSL2809"),
]

LEVELS = {"S", "A", "B", "基础"}
COLLOCATION_SEP = re.compile(r"[、;；\n]+")


def to_int(value) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def to_str(value) -> str:
    return "" if value is None else str(value).strip()


def slugify(word: str) -> str:
    """ASCII-folded slug: ``résumé`` -> ``resume``, ``café`` -> ``cafe``.

    Accents are kept in the displayed headword but must not leak into the id —
    ``r-sum`` / ``caf`` are neither readable nor stable as favourite keys.
    """
    folded = unicodedata.normalize("NFKD", word).encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^a-z0-9]+", "-", folded.lower()).strip("-")
    return slug or "word"


def assign_ids(items: list[dict]) -> None:
    """Give every headword one stable id, unique inside each word list.

    The same headword keeps the same id across lists (ADR-0006: a favourite is
    keyed on the word, not on a list row). Collisions after folding (``resume``
    vs ``résumé``) are resolved in alphabetical order of the raw headword, so
    the numbering does not shift when frequencies or scores change.
    """
    id_by_word: dict[str, str] = {}
    used: set[str] = set()
    for word in sorted({it["word"] for it in items}):
        base = slugify(word)
        candidate, n = base, 1
        while candidate in used:
            n += 1
            candidate = f"{base}-{n}"
        used.add(candidate)
        id_by_word[word] = candidate
    for it in items:
        it["id"] = id_by_word[it["word"]]


def norm_expr(value: str) -> str:
    """Normalize a collocation expression for best-effort matching."""
    s = to_str(value).lower()
    s = s.replace("…", "...").replace("—", "-")
    return re.sub(r"\s+", " ", s).strip()


def tight_expr(value: str) -> str:
    """Aggressive normalization: letters/digits only, no separators."""
    return re.sub(r"[^a-z0-9]+", "", to_str(value).lower())


def index_collocations(path: Path) -> dict[str, str]:
    """Map (normalized expression, tight expression) -> collocation id."""
    if not path.exists():
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    index: dict[str, str] = {}
    for it in data.get("items", []):
        expr = it.get("expression", "")
        cid = it.get("id")
        if not expr or not cid:
            continue
        index[norm_expr(expr)] = cid
        tight = tight_expr(expr)
        if tight:
            index.setdefault(tight, cid)
    return index


def build_items(ws, list_key: str, coll_index: dict[str, str]) -> list[dict]:
    # Map header name -> column index from row 0, so column order drift is tolerated.
    rows = ws.iter_rows(values_only=True)
    header = [to_str(c) for c in next(rows)]
    col = {name: i for i, name in enumerate(header)}

    def cell(row, name):
        i = col.get(name)
        return row[i] if i is not None and i < len(row) else None

    items: list[dict] = []
    for row in rows:
        word = to_str(cell(row, "单词"))
        if not word:
            continue
        coll_label = to_str(cell(row, "关联固定搭配"))
        coll_id = None
        if coll_label:
            # a cell may hold several expressions; match the first that resolves
            for part in COLLOCATION_SEP.split(coll_label):
                part = part.strip()
                if not part:
                    continue
                coll_id = coll_index.get(norm_expr(part)) or coll_index.get(tight_expr(part))
                if coll_id:
                    break
        items.append({
            "id": slugify(word),
            "list": list_key,
            "word": word,
            "phonetic": to_str(cell(row, "音标")) or None,
            "pos": to_str(cell(row, "词性")) or None,
            "zhDef": to_str(cell(row, "中文释义")),
            "enDef": to_str(cell(row, "英文简释")),
            "source": to_str(cell(row, "词表来源")) or None,
            "ngslRank": to_int(cell(row, "NGSL Rank")),
            "tslRank": to_int(cell(row, "TSL Rank")),
            "totalFreq": to_int(cell(row, "题库总次数")) or 0,
            "partCounts": {
                "5": to_int(cell(row, "Part5次数")) or 0,
                "6": to_int(cell(row, "Part6次数")) or 0,
                "7": to_int(cell(row, "Part7次数")) or 0,
            },
            "p5AnswerCount": to_int(cell(row, "Part5答案次数")) or 0,
            "collocation": coll_label or None,
            "collocationCount": to_int(cell(row, "搭配数")) or 0,
            "collocationId": coll_id,
            "wordClassTag": to_str(cell(row, "词类标签")) or None,
            "core1500": to_str(cell(row, "核心1500")) == "是",
            "compositeScore": round(float(cell(row, "综合分") or 0), 2),
            "level": (lambda v: v if v in LEVELS else None)(to_str(cell(row, "推荐级别"))),
            "sourceUrls": {
                "list": to_str(cell(row, "词表来源URL")) or None,
                "def": to_str(cell(row, "中文释义来源URL")) or None,
            },
        })
    return items


def main(argv=None) -> int:
    argv = argv or sys.argv[1:]
    repo = Path(__file__).resolve().parent.parent
    src = repo / "单词本.xlsx"
    out_dir = repo / "data" / "words"
    coll_path = repo / "data" / "collocations" / "collocations.json"
    if "--source" in argv:
        src = Path(argv[argv.index("--source") + 1])
    if "--out" in argv:
        out_dir = Path(argv[argv.index("--out") + 1])

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    coll_index = index_collocations(coll_path)
    if not coll_index:
        print(f"  note: {coll_path.name} not found / empty — collocation links will be empty")

    all_items: list[dict] = []
    lists_meta: dict[str, dict] = {}
    for sheet_name, key, label in SHEETS:
        if sheet_name not in wb.sheetnames:
            raise SystemExit(f"sheet {sheet_name!r} not found in {src.name}; sheets: {wb.sheetnames}")
        items = build_items(wb[sheet_name], key, coll_index)
        # deterministic order: composite score desc, then word — matches the xlsx's own ranking
        items.sort(key=lambda r: (-r["compositeScore"], r["word"].lower()))
        all_items.extend(items)
        lists_meta[key] = {"label": label, "count": len(items)}

    assign_ids(all_items)

    out_dir.mkdir(parents=True, exist_ok=True)
    payload = {
        "source": f"{src.name} · {' + '.join(s for s, _, _ in SHEETS)}",
        "count": len(all_items),
        "lists": lists_meta,
        "items": all_items,
    }
    out_path = out_dir / "words.json"
    out_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    # summary, mirroring extract_collocations.py's reporting style
    print(f"wrote {len(all_items)} words to {out_dir}")
    for key, meta in lists_meta.items():
        print(f"  {meta['label']:<22} {meta['count']}")
    by_level = {"S": 0, "A": 0, "B": 0, "基础": 0, None: 0}
    for it in all_items:
        by_level[it["level"]] += 1
    print(
        "  level: "
        f"S={by_level['S']} A={by_level['A']} B={by_level['B']} 基础={by_level['基础']} none={by_level[None]}"
    )
    linked = sum(1 for it in all_items if it["collocationId"])
    print(f"  collocation links resolved: {linked}")
    flagged = [it["word"] for it in all_items if not it["zhDef"]]
    if flagged:
        print(f"  WARNING: {len(flagged)} words with empty 中文释义: {flagged[:5]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
