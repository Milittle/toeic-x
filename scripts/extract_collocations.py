#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl"]
# ///
"""Extract the TOEIC collocations / idioms workbook into a static JSON library.

Source : ``TOEIC_Reading_固定搭配与惯用法_最终版.xlsx`` (sheet ``高频汇总``)
Output : ``data/collocations/collocations.json``  (read-only, ADR-0004)

The collocations are a static reference corpus parallel to the question bank
(ADR-0002). The app never writes this file; regeneration is manual via
``uv run scripts/extract_collocations.py``.

Columns are read by **header name** (so column order may drift): the workbook's
own 表达 / 中文 / 类型 / 频次 / 优先级 / 来源类型, plus three optional columns
added for 自撰例句 — ``例句`` / ``例句译文`` / ``例句备注`` (only the collocations
with no example in the question bank have them; see
``.scratch/collocation-examples/spec.md``).

Question-source links (``sources``) are NOT trusted from the xlsx — its
self-reported question numbers never matched the PDF-extracted bank. Instead,
after building the records we scan the question bank text to record where each
expression actually appears (see ``collocation_sources.compute_sources``).
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import openpyxl

from collocation_sources import compute_sources

SHEET = "高频汇总"

# Known chat-transcript pollution that leaked into a few 中文 cells during the
# book's PDF extraction (e.g. the "at one's fingertips" row). We cut the cell at
# the first such marker so only the genuine gloss survives.
POLLUTION_MARKERS = [
    "格兰特", "格斯", "约兰达", "彼得", "玛丽", "卡特",
    "[下午", "[上午", "大家好", "还记得", "检查员",
]


def clean_chinese(value) -> str:
    if value is None:
        return ""
    text = str(value).split("\n")[0].strip()  # keep only the first line
    cut = len(text)
    for marker in POLLUTION_MARKERS:
        idx = text.find(marker)
        if idx != -1:
            cut = min(cut, idx)
    return re.sub(r"\s+", " ", text[:cut]).strip()


def to_int(value) -> int:
    if value is None or value == "":
        return 0
    try:
        return int(value)
    except (TypeError, ValueError):
        return 0


def slugify(expression: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", expression.lower()).strip("-")
    return slug or "item"


def build_records(rows, col: dict[str, int]) -> list[dict]:
    """按表头名取列（列顺序漂移不影响结果），对齐 extract_words.py 的做法。"""

    def cell(row, name):
        i = col.get(name)
        return row[i] if i is not None and i < len(row) else None

    records = []
    for row in rows:
        expression = (str(cell(row, "表达")).strip() if cell(row, "表达") else "")
        if not expression:
            continue
        priority = (str(cell(row, "优先级")).strip() if cell(row, "优先级") is not None else "")
        record = {
            "expression": expression,
            "chinese": clean_chinese(cell(row, "中文")),
            "type": (str(cell(row, "类型")).strip() if cell(row, "类型") is not None else ""),
            "bookFreq": to_int(cell(row, "本书词汇栏/解析收录次数")),
            "counts": {
                "5": to_int(cell(row, "P5次数")),
                "6": to_int(cell(row, "P6次数")),
                "7": to_int(cell(row, "P7次数")),
            },
            "bookletCount": to_int(cell(row, "题册原文精确出现次数")),
            "priority": priority if priority in {"S", "A", "B"} else None,
            "sourceType": (
                str(cell(row, "来源类型")).strip() if cell(row, "来源类型") is not None else ""
            ),
            # question links are filled by scanning the bank, not from the xlsx.
            "sources": [],
        }
        # 自撰例句（只给题库里没有用例的搭配补，见 .scratch/collocation-examples/spec.md）：
        # 没有就不写这三个键，避免整库 JSON 平白长大。
        example = str(cell(row, "例句")).strip() if cell(row, "例句") else ""
        if example:
            record["example"] = example
            zh = str(cell(row, "例句译文")).strip() if cell(row, "例句译文") else ""
            if zh:
                record["exampleZh"] = zh
            note = str(cell(row, "例句备注")).strip() if cell(row, "例句备注") else ""
            if note:
                record["exampleNote"] = note
        records.append(record)
    # deterministic order so generated ids are stable across runs
    records.sort(key=lambda r: (r["expression"].lower(), r["chinese"]))
    # assign stable, collision-safe ids
    slug_counts: dict[str, int] = {}
    for rec in records:
        base = slugify(rec["expression"])
        n = slug_counts.get(base, 0) + 1
        slug_counts[base] = n
        rec["id"] = base if n == 1 else f"{base}-{n}"
    return records


def main(argv=None) -> int:
    argv = argv or sys.argv[1:]
    repo = Path(__file__).resolve().parent.parent
    src = repo / "TOEIC_Reading_固定搭配与惯用法_最终版.xlsx"
    out_dir = repo / "data" / "collocations"
    if "--source" in argv:
        src = Path(argv[argv.index("--source") + 1])
    if "--out" in argv:
        out_dir = Path(argv[argv.index("--out") + 1])

    wb = openpyxl.load_workbook(src, data_only=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"sheet {SHEET!r} not found in {src.name}; sheets: {wb.sheetnames}")
    ws = wb[SHEET]
    header = [str(c.value).strip() if c.value is not None else "" for c in next(ws.iter_rows(min_row=1))]
    col = {name: i for i, name in enumerate(header) if name}
    if "表达" not in col:
        raise SystemExit(f"sheet {SHEET!r} 的表头里找不到「表达」列，中止。")
    rows = [tuple(c.value for c in row) for row in ws.iter_rows(min_row=2, values_only=False)]

    items = build_records(rows, col)
    stats = compute_sources(items, repo / "data" / "questions")

    out_dir.mkdir(parents=True, exist_ok=True)
    payload = {
        "source": f"{src.name} · {SHEET}",
        "count": len(items),
        "items": items,
    }
    out_path = out_dir / "collocations.json"
    out_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    # summary, mirroring extract_question_bank.py's reporting style
    by_priority = {"S": 0, "A": 0, "B": 0, None: 0}
    for it in items:
        by_priority[it["priority"]] += 1
    print(f"wrote {len(items)} collocations to {out_dir}")
    print(
        "  priority: "
        f"S={by_priority['S']} A={by_priority['A']} B={by_priority['B']} none={by_priority[None]}"
    )
    print(
        f"  with question links: {stats['matched']}/{len(items)} "
        f"({stats['links']} links total)"
    )
    if stats["zero"]:
        print(f"  zero-link (absent or structurally unmatched): {len(stats['zero'])}")
    flagged = [it["expression"] for it in items if len(it["chinese"]) == 0]
    if flagged:
        print(f"  WARNING: {len(flagged)} rows with empty gloss after cleaning: {flagged[:5]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
