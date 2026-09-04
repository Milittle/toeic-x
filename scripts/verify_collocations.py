#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl"]
# ///
"""全量结构自检搭配库，生成人工核对清单。

只读，不修改 ``data/collocations/collocations.json`` 或 Excel。它把当前搭配库条目
搭配按七类可疑问题归类，输出 ``data/collocations/VERIFICATION.md``，供人
对照纸质书 / Excel 逐条核对。

用法::

    uv run scripts/verify_collocations.py          # 或 python3 scripts/verify_collocations.py

修复入口是 Excel 源文件 ``高频汇总`` 表，改完再 ``npm run collocations``
重新生成 JSON；真题链接问题则由 ``uv run scripts/collocation_sources.py``
重扫。本脚本不提供 ``--apply``——搭配库没有 ``verified`` 标记，正确与否
以人工核对为准。

检查项
------
A. 表达列为碎片/垃圾        —— 含中文/全角标点，或表达式以 ``(`` 结尾（列对齐已错乱）
B. 表达带括号/等号注解      —— 如 ``canteen (= cafeteria)``，需决定是否剥离注解
C. 释义段落泄漏            —— 释义超过阈值，正文被并进了释义
D. 释义碎片/截断           —— 以「表示/和/包括/以及」开头、含 ``；be`` 或过短
E. 近似重复               —— 归一化后相同，或一者归一化后是另一者的前缀
F. 零真题链接             —— 未匹配到任何题目（需人工确认是否真的缺）
G. 链接虚高               —— 关联题目数超过阈值（常见词被过度匹配的假阳性）
"""

from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
COLLOC = REPO / "data" / "collocations" / "collocations.json"
XLSX = REPO / "TOEIC_Reading_固定搭配与惯用法_最终版.xlsx"
OUT = REPO / "data" / "collocations" / "VERIFICATION.md"
SHEET = "高频汇总"

GLOSS_MAX = 30          # 释义超过这个字符数就视为段落泄漏
LINK_HIGH = 40          # 关联题目数超过这个值就进入假阳性排查
PREFIX_PAIR_LIMIT = 80  # 前缀近似重复最多展示多少对，避免清单过长

CJK = re.compile(r"[\u4e00-\u9fff]")
FULLWIDTH = re.compile(r"[《》（）＝、，。“”＆…]")
TRAILING_OPEN = re.compile(r"\($")
PAREN_EQ = re.compile(r"[()=]")
BROKEN_START = re.compile(r"^(表示|和|包括|以及|的结构|答案是)")
BROKEN_MID = re.compile(r"；be")


def load_items() -> list[dict]:
    payload = json.loads(COLLOC.read_text(encoding="utf-8"))
    return payload["items"]


def norm(expr: str) -> str:
    """归一化表达式：只保留小写字母数字，用于找重复/前缀关系。"""
    s = expr.lower().replace("’", "'")
    return re.sub(r"[^a-z0-9]+", "", s)


def excel_row_map() -> dict[str, int]:
    """表达式 -> Excel 行号，便于直接回到源文件定位。"""
    rows: dict[str, int] = {}
    try:
        import openpyxl
    except ImportError:
        return rows
    try:
        wb = openpyxl.load_workbook(XLSX, data_only=True, read_only=True)
    except FileNotFoundError:
        return rows
    if SHEET in wb.sheetnames:
        for i, r in enumerate(wb[SHEET].iter_rows(min_row=2, values_only=True), start=2):
            if r and r[0] is not None:
                rows[str(r[0]).strip()] = i
    return rows


def trunc(text: str, n: int) -> str:
    text = re.sub(r"\s+", " ", text)
    return text if len(text) <= n else text[: n - 1] + "…"


def table(lines: list[str], items: list[dict], rows: dict[str, int]) -> None:
    if not items:
        lines.append("_无_")
        lines.append("")
        return
    lines.append("| 表达 | 释义 | 类型 | 优先级 | 链接 | Excel 行 |")
    lines.append("|---|---|---|---|---|---|")
    for it in items:
        expr = it["expression"]
        chinese = trunc(it["chinese"], 34)
        typ = it["type"] or "—"
        pri = it["priority"] or "—"
        nlink = len(it["sources"])
        xrow = rows.get(expr, "")
        lines.append(
            f"| `{expr}` | {chinese} | {typ} | {pri} | {nlink} | {xrow} |"
        )
    lines.append("")


def build_report() -> str:
    items = load_items()
    rows = excel_row_map()

    a = [it for it in items if CJK.search(it["expression"]) or FULLWIDTH.search(it["expression"]) or TRAILING_OPEN.search(it["expression"])]
    b = [it for it in items if not CJK.search(it["expression"]) and PAREN_EQ.search(it["expression"])]
    c = [it for it in items if len(it["chinese"]) > GLOSS_MAX]
    d = [it for it in items if BROKEN_START.search(it["chinese"]) or BROKEN_MID.search(it["chinese"]) or len(it["chinese"].strip()) <= 1]
    f = [it for it in items if not it["sources"]]
    g = [it for it in items if len(it["sources"]) > LINK_HIGH]
    missing = [
        it for it in items
        if not it["chinese"].strip() or not it["type"].strip() or not it["priority"]
    ]

    # E: near duplicates — exact normalized equal + prefix relations
    by_norm: dict[str, list[dict]] = defaultdict(list)
    for it in items:
        by_norm[norm(it["expression"])].append(it)
    exact_dups = sorted(
        (g for g in by_norm.values() if len({x["expression"] for x in g}) > 1),
        key=lambda g: g[0]["expression"].lower(),
    )
    prefix_pairs: list[tuple[dict, dict]] = []
    norms = sorted(by_norm, key=len)
    for i, na in enumerate(norms):
        for nb in norms[i + 1 :]:
            if nb.startswith(na) and len(nb) > len(na):
                prefix_pairs.append((by_norm[na][0], by_norm[nb][0]))
            if len(prefix_pairs) >= PREFIX_PAIR_LIMIT:
                break
        if len(prefix_pairs) >= PREFIX_PAIR_LIMIT:
            break

    lines: list[str] = []
    lines.append("# 搭配库校验清单")
    lines.append("")
    lines.append("> 只读自检，逐类列出可疑条目供人工核对。核对后修改 Excel `高频汇总`，")
    lines.append("> 再 `npm run collocations` 重新生成；真题链接用 `uv run scripts/collocation_sources.py` 重扫。")
    lines.append("")
    lines.append("| 检查项 | 数量 |")
    lines.append("|---|---|")
    lines.append(f"| A 表达列为碎片/垃圾 | {len(a)} |")
    lines.append(f"| B 表达带括号/等号注解 | {len(b)} |")
    lines.append(f"| C 释义段落泄漏 | {len(c)} |")
    lines.append(f"| D 释义碎片/截断 | {len(d)} |")
    lines.append(f"| E 近似重复（精确归一化） | {len(exact_dups)} |")
    lines.append(f"| E 近似重复（前缀关系，展示 ≤{PREFIX_PAIR_LIMIT}） | {len(prefix_pairs)} |")
    lines.append(f"| F 零真题链接 | {len(f)} |")
    lines.append(f"| G 链接虚高 (>{LINK_HIGH}) | {len(g)} |")
    lines.append(f"| 字段缺失（释义/类型/优先级为空） | {len(missing)} |")
    lines.append("")
    lines.append(f"总条目：{len(items)}")
    lines.append("")

    lines.append("## A. 表达列为碎片/垃圾（列对齐错乱，需删除或修复）")
    lines.append("")
    table(lines, a, rows)

    lines.append("## B. 表达带括号/等号注解（决定是否剥离「见另」注解）")
    lines.append("")
    table(lines, b, rows)

    lines.append(f"## C. 释义段落泄漏（正文被并进释义，>{GLOSS_MAX} 字）")
    lines.append("")
    table(lines, c, rows)

    lines.append("## D. 释义碎片/截断（以「表示/和/包括/以及」开头、含「；be」或过短）")
    lines.append("")
    table(lines, d, rows)

    lines.append("## E. 近似重复")
    lines.append("")
    lines.append("### 精确归一化后相同（仅空格/标点差异）")
    lines.append("")
    if not exact_dups:
        lines.append("_无_")
        lines.append("")
    for group in exact_dups:
        lines.append("  " + "　".join(f"`{x['expression']}`（{trunc(x['chinese'], 20)}）" for x in group))
        lines.append("")
    lines.append("### 前缀关系（一者是另一者的省略形式，人工判断是否保留两条）")
    lines.append("")
    if not prefix_pairs:
        lines.append("_无_")
        lines.append("")
    for x, y in prefix_pairs:
        lines.append(f"- `{x['expression']}` ⊂ `{y['expression']}`（{trunc(x['chinese'], 16)} / {trunc(y['chinese'], 16)}）")
    lines.append("")

    lines.append("## F. 零真题链接（未匹配到任何题目，需人工确认）")
    lines.append("")
    table(lines, f, rows)

    lines.append(f"## G. 链接虚高（>{LINK_HIGH} 处，常见词过度匹配的假阳性排查）")
    lines.append("")
    table(lines, g, rows)

    if missing:
        lines.append("## 字段缺失")
        lines.append("")
        table(lines, missing, rows)

    return "\n".join(lines)


def main(argv=None) -> int:
    report = build_report()
    OUT.write_text(report, encoding="utf-8")
    print(f"wrote {OUT}")
    # 打印顶部汇总表，方便快速判断规模
    printing = False
    for line in report.splitlines():
        if line.startswith("| 检查项"):
            printing = True
            continue
        if not printing:
            continue
        if line.startswith("|---"):
            continue
        if line.startswith("|"):
            label, _, num = line.strip("|").partition("|")
            print(f"  {label.strip()}: {num.strip()}")
        else:
            break
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
