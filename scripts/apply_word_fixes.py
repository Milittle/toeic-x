#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl"]
# ///
"""把审计结果回写进词库的 Excel 源文件 ``单词本.xlsx``。

审计（结构自检 + LLM 释义审查）产出的是「修复计划」——一个 JSON 文件，列出对
``重点1500 / TOEIC专项_TSL1250 / 基础_NGSL2809`` 三张表的操作。本脚本只负责
**按计划改 Excel**，不重生成 JSON（改完再 ``npm run words``）。

修复计划格式（``data/words/fix_plan.json``）::

    {
      "note": "人工确认后的修复计划",
      "ops": [
        {"op": "set_gloss",    "word": "refund", "gloss": "n. 退款；vt. 退还，退款"},
        {"op": "set_phonetic", "word": "laptop", "phonetic": "'læptɒp"},
        {"op": "set_endef",    "word": "hotline", "endef": "n. a direct telephone line"}
      ]
    }

- ``set_gloss``     改「中文释义」列；
- ``set_phonetic``  改「音标」列；
- ``set_endef``     改「英文简释」列。

同一个单词可能同时出现在多张表里（``words.json`` 里字段完全一致），默认**三张表
一起改**；需要只改某张表时加 ``"lists": ["key1500"]``。``word`` 与「单词」列需逐字
匹配。

写入前会自动备份成 ``单词本.bak-YYYYmmdd-HHMMSS.xlsx``；``--dry-run`` 只打印。

用法::

    uv run scripts/apply_word_fixes.py --from-llm            # 从 llm_review.jsonl 生成计划
    uv run scripts/apply_word_fixes.py --from-llm \
        --review data/words/llm_review_round2.jsonl --out-plan data/words/llm_fix_plan_round2.json
    uv run scripts/apply_word_fixes.py --plan data/words/fix_plan.json --dry-run
    uv run scripts/apply_word_fixes.py --plan data/words/fix_plan.json
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import time
import zipfile
from collections import Counter
from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl

REPO = Path(__file__).resolve().parent.parent
XLSX = REPO / "单词本.xlsx"
REVIEW_JSONL = REPO / "data" / "words" / "llm_review.jsonl"
PLAN = REPO / "data" / "words" / "fix_plan.json"
LLM_PLAN = REPO / "data" / "words" / "llm_fix_plan.json"
REJECT_MD = REPO / "data" / "words" / "QA_REJECTS.md"
FIX_LOG = REPO / "data" / "words" / "FIX_LOG.md"

SHEET_FOR_LIST = {
    "key1500": "重点1500",
    "tsl1250": "TOEIC专项_TSL1250",
    "ngsl2809": "基础_NGSL2809",
}

# 列号（1 起），见 extract_words.py 读表头的方式
COL_WORD = 2
COL_PHONETIC = 3
COL_ZH_DEF = 5
COL_EN_DEF = 6
COL_NOTE = 22

COLUMN_FOR_OP = {
    "set_gloss": COL_ZH_DEF,
    "set_phonetic": COL_PHONETIC,
    "set_endef": COL_EN_DEF,
    "set_note": COL_NOTE,
}
COL_LETTER = {2: "B", 3: "C", 5: "E", 6: "F", 22: "V"}
FIELD_FOR_OP = {
    "set_gloss": "gloss",
    "set_phonetic": "phonetic",
    "set_endef": "endef",
    "set_note": "note",
}


def row_map(ws) -> dict[str, int]:
    rows: dict[str, int] = {}
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if r and len(r) >= COL_WORD and r[COL_WORD - 1] is not None:
            rows[str(r[COL_WORD - 1]).strip()] = i
    return rows


CJK = re.compile(r"[\u4e00-\u9fff]")
POS_PREFIX = re.compile(r"^\s*[a-z]{1,7}\.")
MARKDOWN = re.compile(r"[|*`#]|\]\(|\\[rnt]")
GLOSS_MAX = 60


def qa_reject(word: str, gloss: str, current: str) -> str | None:
    """LLM 建议的机械质检：拦掉明显不能直接进 Excel 的建议，返回拒绝原因。"""
    if not gloss.strip():
        return "空释义"
    if gloss.strip() == (current or "").strip():
        return "与原释义相同（no-op）"
    if not CJK.search(gloss):
        return "没有中文"
    if POS_PREFIX.match(current or "") and not POS_PREFIX.match(gloss):
        return "丢失词性前缀"
    if len(gloss) > GLOSS_MAX and len(gloss) > len(current or "") * 0.7:
        return f"新释义仍超过 {GLOSS_MAX} 字且压缩不明显"
    if MARKDOWN.search(gloss):
        return "含 Markdown/转义噪声"


def from_llm(plan_path: Path, review_path: Path = REVIEW_JSONL) -> None:
    if not review_path.exists():
        print(f"找不到 {review_path}，请先跑 review_words_llm.py", file=sys.stderr)
        raise SystemExit(2)
    ops, rejects, seen = [], [], {}
    for line in review_path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        obj = json.loads(line)
        if obj.get("ok") is not False or not obj.get("gloss"):
            continue
        seen[obj["word"]] = obj  # 同一单词重复审查时以最后一次为准
    for word, obj in seen.items():
        reason = qa_reject(word, obj["gloss"], obj.get("current_gloss", ""))
        if reason:
            rejects.append((word, obj["gloss"], reason))
            continue
        ops.append({
            "op": "set_gloss",
            "word": word,
            "gloss": obj["gloss"],
            "reason": obj.get("reason", ""),
        })
    plan = {
        "note": "由 llm_review.jsonl 的 ok=false 条目经 QA 过滤后生成，人工确认后再 --plan 应用。",
        "ops": ops,
    }
    plan_path.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"wrote {plan_path}：{len(ops)} 条 set_gloss")
    if rejects:
        lines = ["# 被 QA 拦下的 LLM 建议", "", "| 单词 | 建议释义 | 原因 |", "|---|---|---|"]
        lines += [f"| `{w}` | {g} | {r} |" for w, g, r in sorted(rejects, key=lambda x: x[0].lower())]
        REJECT_MD.write_text("\n".join(lines) + "\n", encoding="utf-8")
        print(f"QA 拦下 {len(rejects)} 条，见 {REJECT_MD}")
    print("请复核后再用 --plan 应用。")


def apply_plan(plan_path: Path, dry_run: bool) -> None:
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    ops = plan.get("ops", [])
    if not ops:
        print("计划为空，无事可做。")
        return

    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    maps = {key: row_map(wb[sheet]) for key, sheet in SHEET_FOR_LIST.items() if sheet in wb.sheetnames}
    wb.close()

    matched, missing = 0, []
    touched_sheets: set[str] = set()
    edits: dict[str, dict[str, str]] = {}

    for op in ops:
        kind = op.get("op")
        word = str(op.get("word", "")).strip()
        field = FIELD_FOR_OP.get(kind)
        if field is None:
            missing.append(f"{word} (未知 op: {kind!r})")
            continue
        value = op.get(field, "")
        lists = op.get("lists") or list(maps)
        hits = 0
        for key in lists:
            sheet = SHEET_FOR_LIST.get(key)
            if sheet is None or key not in maps:
                missing.append(f"{word} (未知词表: {key!r})")
                continue
            row = maps[key].get(word)
            if row is None:
                continue
            hits += 1
            touched_sheets.add(sheet)
            print(f"  {kind:12} {sheet} row {row}: {word!r} -> {value!r}")
            if not dry_run:
                edits.setdefault(sheet, {})[f"{COL_LETTER[COLUMN_FOR_OP[kind]]}{row}"] = value
        if hits:
            matched += 1
        else:
            missing.append(f"{word}（在 {'/'.join(lists)} 中未找到）")

    print(f"\n计划 {len(ops)} 条：命中单词 {matched}，未命中/跳过 {len(missing)}")
    for m in missing:
        print(f"  未命中: {m}")

    if dry_run:
        print("dry-run：未写入 Excel。")
        return

    backup = XLSX.with_name(XLSX.stem + f".bak-{time.strftime('%Y%m%d-%H%M%S')}.xlsx")
    shutil.copy2(XLSX, backup)
    written = patch_xlsx(XLSX, edits)
    print(f"已备份到 {backup.name}")
    print(f"已写入 {XLSX.name}（{len(touched_sheets)} 张表，{written} 个单元格）")

    stamps = time.strftime("%Y-%m-%d %H:%M:%S")
    kinds = Counter(op.get("op") for op in ops)
    entry = [
        f"## {stamps}",
        "",
        f"- 计划：`{plan_path.name}`（{plan.get('note', '')}）",
        f"- 应用：{'、'.join(f'{k} {v} 条' for k, v in kinds.most_common())}；命中单词 {matched}，未命中 {len(missing)}",
        f"- 写入：{len(touched_sheets)} 张表 / {written} 个单元格；备份 `{backup.name}`",
        "",
    ]
    header = "" if FIX_LOG.exists() else "# 词库修复记录\n\n> 由 `scripts/apply_word_fixes.py` 追加；改完 Excel 仍需 `npm run words` 重新生成 JSON。\n\n"
    with FIX_LOG.open("a", encoding="utf-8") as f:
        f.write(header + "\n".join(entry))
    print(f"已追加记录到 {FIX_LOG}")


def sheet_xml_paths(zf: zipfile.ZipFile) -> dict[str, str]:
    """工作表名 -> xl/worksheets/sheetN.xml（经 workbook.xml 的 r:id 与 rels 解析）。"""
    wb_xml = zf.read("xl/workbook.xml").decode("utf-8")
    rels_xml = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rid2target: dict[str, str] = {}
    for tag in re.findall(r"<Relationship\b[^>]*/?>", rels_xml):
        rid = re.search(r'Id="([^"]+)"', tag)
        target = re.search(r'Target="([^"]+)"', tag)
        if rid and target:
            rid2target[rid.group(1)] = target.group(1)
    paths: dict[str, str] = {}
    for tag in re.findall(r"<(?:[A-Za-z0-9]+:)?sheet\b[^>]*/?>", wb_xml):
        name = re.search(r'name="([^"]*)"', tag)
        rid = re.search(r'r:id="([^"]+)"', tag)
        if not name or not rid:
            continue
        target = rid2target.get(rid.group(1), "")
        target = target.lstrip("/")
        if target and not target.startswith("xl/"):
            target = "xl/" + target
        if target:
            paths[name.group(1)] = target
    return paths


def patch_sheet_xml(xml: str, edits: dict[str, str], prefix: str) -> tuple[str, list[str]]:
    """把指定单元格改成内联字符串，其余字节原样保留（含公式的缓存值）。"""
    p = f"{prefix}:" if prefix else ""
    missing: list[str] = []
    for ref, value in edits.items():
        pat = re.compile(
            rf"<{p}c\b[^>]*\br=\"{re.escape(ref)}\"[^>]*?(?:/>|>.*?</{p}c>)",
            re.S,
        )
        m = pat.search(xml)
        if not m:
            missing.append(ref)
            continue
        tag = m.group(0)
        attrs = re.match(rf"<{p}c\b([^>]*?)(?:/>|>)", tag, re.S).group(1)
        attrs = re.sub(r'\s+t="[^"]*"', "", attrs).rstrip()
        body = (
            f"<{p}is><{p}t xml:space=\"preserve\">{escape(value)}</{p}t></{p}is>"
        )
        xml = xml[: m.start()] + f'<{p}c{attrs} t="inlineStr">{body}</{p}c>' + xml[m.end():]
    return xml, missing


def patch_xlsx(path: Path, edits: dict[str, dict[str, str]]) -> int:
    """按 {工作表: {单元格: 新值}} 直接改写 xlsx 包内的 sheet XML。

    为什么不用 openpyxl 保存：``单词本.xlsx`` 的 ``序号`` / ``综合分`` / ``推荐级别``
    是公式，openpyxl 重新保存会丢掉公式的缓存值，``data_only=True`` 再读就全是
    None（``level`` 和 ``compositeScore`` 会整列变空）。这里只替换目标单元格，
    公式、缓存值、样式、表格定义都原样保留。
    """
    tmp = path.with_suffix(".xlsx.tmp")
    written = 0
    with zipfile.ZipFile(path) as zin:
        paths = sheet_xml_paths(zin)
        targets = {paths[name]: edits_for for name, edits_for in edits.items() if name in paths}
        unknown = [name for name in edits if name not in paths]
        if unknown:
            raise SystemExit(f"找不到工作表: {unknown}")
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename in targets:
                    xml = data.decode("utf-8")
                    prefix_match = re.match(r"<\?xml[^>]*\?>\s*<([A-Za-z0-9]+):worksheet", xml)
                    prefix = prefix_match.group(1) if prefix_match else ""
                    xml, missing = patch_sheet_xml(xml, targets[item.filename], prefix)
                    if missing:
                        raise SystemExit(f"{item.filename} 中找不到单元格: {missing[:5]}")
                    written += len(targets[item.filename])
                    data = xml.encode("utf-8")
                zout.writestr(item, data)
    tmp.replace(path)
    return written


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--from-llm", action="store_true",
                        help="从 llm_review.jsonl 生成 llm_fix_plan.json")
    parser.add_argument("--plan", type=Path, default=PLAN, help="修复计划 JSON")
    parser.add_argument("--dry-run", action="store_true", help="只打印，不写 Excel")
    parser.add_argument("--review", type=Path, default=REVIEW_JSONL,
                        help="审查结果 JSONL（复审用 data/words/llm_review_round2.jsonl）")
    parser.add_argument("--out-plan", type=Path, default=LLM_PLAN,
                        help="--from-llm 生成的计划写到哪（复审用 llm_fix_plan_round2.json）")
    args = parser.parse_args(argv)

    if args.from_llm:
        from_llm(args.out_plan, args.review)
        return 0
    if not args.plan.exists():
        print(f"找不到计划文件 {args.plan}。先用 --from-llm 生成，或手工创建。", file=sys.stderr)
        return 2
    apply_plan(args.plan, args.dry_run)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
